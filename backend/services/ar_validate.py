"""
AR validation: ECC registry (ECC_DATA/*.xlsx) vs filled S/4 template
(S4_DATA/*.xlsx).

Reads:
  - ECC: raw extract written by GET /ar/ecc_data
  - S/4: the LTMC "Customer Open Items" workbook written by POST /ar/process

The ECC extract carries human-readable column headers while the checks
below use canonical names; _rename_ecc_columns() bridges the two in one
place using ECC_COLUMN_ALIASES from ar_field_map.
"""
from typing import Any, Callable, Dict, Iterable, List, Optional

import pandas as pd
import openpyxl

from backend.models import mappings
from backend.models.ar_field_map import ECC_COLUMN_ALIASES, S4_FIELDS
from backend.services.ar_processor import _detect_data_start_row, get_tax_code
from backend.utils.format_utils import clean_string, normalize_series


# ============================================================
# AR Validation Configuration
# ============================================================

S4_SHEET_NAME = "Customer Open Items"
TECHNICAL_NAME_ROW = 5

COMPANY_CODE_MAPPING = mappings.COMPANY_CODE_MAPPING


# ============================================================
# Payment Terms Mapping (ECC -> S/4)
# ============================================================
# Sourced from mappings.AR_PAYMENT_TERMS_MAPPING (shared with ar_processor).
#
# The first character of the S/4 term determines the group:
#   N -> Net terms
#   P -> Proxy terms
#   Z -> Discount terms
#   E -> E_payment_terms (special group)
# ============================================================

_NORMALIZED_PAYMENT_MAPPING = {
    key.lstrip("0"): value for key, value in mappings.AR_PAYMENT_TERMS_MAPPING.items()
}


# ============================================================
# Utility Functions
# ============================================================

def is_non_empty(value: Any) -> bool:
    """Returns True when an Excel cell contains an actual value."""
    if value is None:
        return False
    if isinstance(value, str):
        return value.strip() != ""
    try:
        if pd.isna(value):
            return False
    except (TypeError, ValueError):
        pass
    return True


# ============================================================
# Result-Building Helpers
# ============================================================

def pass_fail(is_match: bool) -> str:
    """Map a boolean match into the "PASS" / "FAIL" status string."""
    return "PASS" if is_match else "FAIL"


def make_detail(
    label: str,
    left_count: Optional[float],
    right_count: Optional[float],
    message: str,
    status: Optional[str] = None,
    money: bool = False,
    **extra: Any,
) -> Dict[str, Any]:
    """Build one row of a check's `details` list."""
    if status is None:
        status = pass_fail(left_count == right_count)

    detail: Dict[str, Any] = {
        "label": label,
        "left_count": left_count,
        "right_count": right_count,
        "status": status,
        "message": message,
    }
    if money:
        detail["money"] = True
    detail.update(extra)
    return detail


def make_check(
    check_name: str,
    details: List[Dict[str, Any]],
    passing_message: str,
    failing_message: str,
    **extra: Any,
) -> Dict[str, Any]:
    """Build a check's top-level result from its already-built `details`."""
    overall_pass = all(detail["status"] == "PASS" for detail in details)
    check: Dict[str, Any] = {
        "check_name": check_name,
        "status": pass_fail(overall_pass),
        "details": details,
        "message": passing_message if overall_pass else failing_message,
    }
    check.update(extra)
    return check


def missing_column_check(check_name: str, source_label: str, column: str) -> Dict[str, Any]:
    """Short-circuit result for a check whose required column is absent."""
    return {
        "check_name": check_name,
        "status": "FAIL",
        "details": [],
        "message": f'{source_label} does not contain a "{column}" column.',
    }


def first_missing_column(df: pd.DataFrame, columns: Iterable[str]) -> Optional[str]:
    """Return the first of `columns` not present in `df`, or None."""
    for column in columns:
        if column not in df.columns:
            return column
    return None


# ============================================================
# Excel Reading
# ============================================================

def _rename_ecc_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Apply ECC_COLUMN_ALIASES (case-insensitive) to the ECC dataframe."""
    rename_map = {}
    for column in df.columns:
        key = str(column).strip().lower()
        canonical = ECC_COLUMN_ALIASES.get(key)
        if canonical:
            rename_map[column] = canonical
    return df.rename(columns=rename_map)


def read_ecc_registry(registry_file) -> pd.DataFrame:
    """Read the ECC AR registry, drop empty rows, alias column names."""
    df = pd.read_excel(registry_file, header=0)
    df = df.dropna(how="all")
    return _rename_ecc_columns(df)


def read_s4_customer_open_items(filled_file) -> pd.DataFrame:
    """
    Read the S/4 Customer Open Items sheet.

    Row 5 carries the technical field names (BUKRS, KUNNR, ...). Data
    start row is detected the same way the processor detects it, so the
    reader and writer agree even if the template layout changes.
    """
    wb = openpyxl.load_workbook(filled_file, read_only=True, data_only=True)

    if S4_SHEET_NAME not in wb.sheetnames:
        raise ValueError(
            f'S/4 file does not contain the required sheet "{S4_SHEET_NAME}".'
        )

    ws = wb[S4_SHEET_NAME]

    technical_headers: List[str] = []
    technical_columns: Dict[str, int] = {}
    for col_index, cell in enumerate(ws[TECHNICAL_NAME_ROW], start=1):
        name = clean_string(cell.value)
        technical_headers.append(name)
        if name:
            technical_columns[name] = col_index

    data_start_row = _detect_data_start_row(ws, technical_columns)

    records = []
    for row in ws.iter_rows(min_row=data_start_row, values_only=True):
        if not any(is_non_empty(value) for value in row):
            continue
        record = {}
        for index, value in enumerate(row):
            if index < len(technical_headers):
                field = technical_headers[index]
                if field:
                    record[field] = value
        records.append(record)

    wb.close()
    return pd.DataFrame(records)


# ============================================================
# Check 1 - Company Code Distribution
# ============================================================

def validate_company_code_counts(ecc_df: pd.DataFrame, s4_df: pd.DataFrame) -> Dict[str, Any]:
    """Per-company-code document counts: ECC code -> mapped S/4 code."""
    missing = first_missing_column(ecc_df, ["Company Code"])
    if missing:
        return missing_column_check("Company Code Distribution", "ECC registry", missing)
    missing = first_missing_column(s4_df, [S4_FIELDS["COMPANY_CODE"]])
    if missing:
        return missing_column_check("Company Code Distribution", "S/4 file", missing)

    ecc_codes = normalize_series(ecc_df["Company Code"])
    s4_codes = normalize_series(s4_df[S4_FIELDS["COMPANY_CODE"]])

    details = []
    for ecc_code in ecc_codes.replace("", pd.NA).dropna().unique():
        ecc_code = ecc_code.upper()
        ecc_count = int((ecc_codes == ecc_code).sum())

        s4_code = COMPANY_CODE_MAPPING.get(ecc_code)
        if not s4_code:
            details.append(make_detail(
                label=f"{ecc_code} (no mapping)",
                left_count=ecc_count, right_count=None,
                status="MAPPING_ERROR",
                message=f"No ECC -> S/4 company code mapping exists for {ecc_code}.",
                ecc_code=ecc_code, s4_code=None,
            ))
            continue

        s4_code = clean_string(s4_code).upper()
        s4_count = int((s4_codes == s4_code).sum())
        details.append(make_detail(
            label=f"{ecc_code} \u2192 {s4_code}",
            left_count=ecc_count, right_count=s4_count,
            message=("Company code count matches."
                     if ecc_count == s4_count
                     else f"Company code count mismatch: ECC={ecc_count}, S/4={s4_count}."),
            ecc_code=ecc_code, s4_code=s4_code,
        ))

    return make_check(
        check_name="Company Code Distribution", details=details,
        passing_message="All company code counts match.",
        failing_message="One or more company code counts do not match.",
    )


# ============================================================
# Check 2 - Customer Distribution
# ============================================================

def validate_customer_distribution(ecc_df: pd.DataFrame, s4_df: pd.DataFrame) -> Dict[str, Any]:
    """Per-customer document counts: ECC customer -> mapped S/4 BP."""
    missing = first_missing_column(ecc_df, ["Customer"])
    if missing:
        return missing_column_check("Customer Distribution", "ECC registry", missing)
    missing = first_missing_column(s4_df, [S4_FIELDS["CUSTOMER"]])
    if missing:
        return missing_column_check("Customer Distribution", "S/4 file", missing)

    ecc_customers = normalize_series(ecc_df["Customer"])
    s4_customers = normalize_series(s4_df[S4_FIELDS["CUSTOMER"]])

    details = []
    for ecc_cust in ecc_customers.replace("", pd.NA).dropna().unique():
        s4_cust = clean_string(mappings.get_s4_ar_customer(ecc_cust))
        ecc_count = int((ecc_customers == ecc_cust).sum())

        if not s4_cust:
            details.append(make_detail(
                label=f"{ecc_cust} (no mapping)",
                left_count=ecc_count, right_count=None,
                status="MAPPING_ERROR",
                message=f"No S/4 business partner mapping for ECC customer {ecc_cust}.",
                ecc_customer=ecc_cust, s4_customer=None,
            ))
            continue

        s4_count = int((s4_customers == s4_cust).sum())
        details.append(make_detail(
            label=f"{ecc_cust} \u2192 {s4_cust}",
            left_count=ecc_count, right_count=s4_count,
            message=("Customer document count matches."
                     if ecc_count == s4_count
                     else f"Customer count mismatch: ECC={ecc_count}, S/4={s4_count}."),
            ecc_customer=ecc_cust, s4_customer=s4_cust,
        ))

    return make_check(
        check_name="Customer Distribution", details=details,
        passing_message="All customer document counts match.",
        failing_message="One or more customer document counts do not match.",
    )


# ============================================================
# Check 3 - Amount Sign Validation
# ============================================================

def validate_sign(ecc_df: pd.DataFrame, s4_df: pd.DataFrame) -> Dict[str, Any]:
    """Validate S/H amount totals company-code-wise."""
    missing = first_missing_column(ecc_df, ["Company Code", "Amount", "Debit/Credit Ind."])
    if missing:
        return missing_column_check("Amount Sign Validation", "ECC registry", missing)
    missing = first_missing_column(s4_df, [S4_FIELDS["COMPANY_CODE"], S4_FIELDS["AMOUNT"]])
    if missing:
        return missing_column_check("Amount Sign Validation", "S/4 file", missing)

    ecc_codes = normalize_series(ecc_df["Company Code"])
    s4_codes = normalize_series(s4_df[S4_FIELDS["COMPANY_CODE"]])
    ecc_amounts = pd.to_numeric(ecc_df["Amount"], errors="coerce").fillna(0)
    ecc_indicator = normalize_series(ecc_df["Debit/Credit Ind."])
    s4_amounts = pd.to_numeric(s4_df[S4_FIELDS["AMOUNT"]], errors="coerce").fillna(0)

    TOLERANCE = 0.01
    details = []

    for ecc_code in ecc_codes.replace("", pd.NA).dropna().unique():
        ecc_code = ecc_code.upper()
        s4_code = COMPANY_CODE_MAPPING.get(ecc_code)

        if not s4_code:
            for direction, label in [("S", "S / Debit (positive)"), ("H", "H / Credit (negative)")]:
                details.append(make_detail(
                    label=f"{ecc_code} (no mapping) | {label}",
                    left_count=None, right_count=None, status="MAPPING_ERROR",
                    message=f"No ECC -> S/4 company code mapping exists for {ecc_code}.",
                    money=True, ecc_code=ecc_code, s4_code=None, direction=direction,
                ))
            continue

        s4_code = clean_string(s4_code).upper()
        ecc_mask = ecc_codes == ecc_code
        s4_mask = s4_codes == s4_code
        ecc_s_total = float(ecc_amounts[ecc_mask & (ecc_indicator == "S")].abs().sum())
        ecc_h_total = float(ecc_amounts[ecc_mask & (ecc_indicator == "H")].abs().sum())
        s4_positive_total = float(s4_amounts[s4_mask & (s4_amounts > 0)].sum())
        s4_negative_total = float(s4_amounts[s4_mask & (s4_amounts < 0)].abs().sum())
        s_match = abs(ecc_s_total - s4_positive_total) < TOLERANCE
        h_match = abs(ecc_h_total - s4_negative_total) < TOLERANCE

        details.append(make_detail(
            label=f"{ecc_code} -> {s4_code} | S / Debit (positive)",
            left_count=ecc_s_total, right_count=s4_positive_total, status=pass_fail(s_match),
            message=("S (debit) total matches S/4 positive WRBTR total." if s_match
                     else f"S (debit) mismatch: ECC={ecc_s_total:.2f}, S/4 positive={s4_positive_total:.2f}."),
            money=True, ecc_code=ecc_code, s4_code=s4_code, direction="S",
        ))
        details.append(make_detail(
            label=f"{ecc_code} -> {s4_code} | H / Credit (negative)",
            left_count=ecc_h_total, right_count=s4_negative_total, status=pass_fail(h_match),
            message=("H (credit) total matches S/4 negative WRBTR total." if h_match
                     else f"H (credit) mismatch: ECC={ecc_h_total:.2f}, S/4 negative={s4_negative_total:.2f}."),
            money=True, ecc_code=ecc_code, s4_code=s4_code, direction="H",
        ))

    return make_check(
        check_name="Amount Sign Validation", details=details,
        passing_message="ECC S/H amount totals match S/4 positive/negative WRBTR totals.",
        failing_message="Amount sign validation failed for one or more company codes.",
        tolerance=TOLERANCE, company_code_wise=True,
    )


# ============================================================
# Check 4 - Total Amount Reconciliation
# ============================================================

def validate_total_amount_reconciliation(ecc_df: pd.DataFrame, s4_df: pd.DataFrame) -> Dict[str, Any]:
    """Grand-total sum of ECC amounts equals grand-total S/4 WRBTR."""
    missing = first_missing_column(ecc_df, ["Amount"])
    if missing:
        return missing_column_check("Total Amount Reconciliation", "ECC registry", missing)
    missing = first_missing_column(s4_df, [S4_FIELDS["AMOUNT"]])
    if missing:
        return missing_column_check("Total Amount Reconciliation", "S/4 file", missing)

    ecc_total = float(pd.to_numeric(ecc_df["Amount"], errors="coerce").fillna(0).abs().sum())
    s4_total = float(pd.to_numeric(s4_df[S4_FIELDS["AMOUNT"]], errors="coerce").fillna(0).abs().sum())
    tolerance = 0.01
    is_match = abs(ecc_total - s4_total) < tolerance

    details = [make_detail(
        label="Grand total |Amount|",
        left_count=ecc_total, right_count=s4_total,
        status=pass_fail(is_match),
        message=("Total amount matches." if is_match
                 else f"Total mismatch: ECC={ecc_total:.2f}, S/4={s4_total:.2f}."),
        money=True,
    )]

    return make_check(
        check_name="Total Amount Reconciliation", details=details,
        passing_message="Sum of ECC amounts equals sum of S/4 WRBTR.",
        failing_message="Sum of ECC amounts does not equal sum of S/4 WRBTR.",
        tolerance=tolerance,
    )


# ============================================================
# Check 5 - Unique Document Number Count
# ============================================================

def validate_unique_document_number_counts(ecc_df: pd.DataFrame, s4_df: pd.DataFrame) -> Dict[str, Any]:
    """Compare unique ECC Document Number vs S/4 XREF1 counts by company code."""
    missing = first_missing_column(ecc_df, ["Company Code", "Document Number"])
    if missing:
        return missing_column_check("Unique Document Number Count", "ECC registry", missing)
    missing = first_missing_column(s4_df, [S4_FIELDS["COMPANY_CODE"], S4_FIELDS["REFERENCE_KEY_1"]])
    if missing:
        return missing_column_check("Unique Document Number Count", "S/4 file", missing)

    ecc_codes = normalize_series(ecc_df["Company Code"])
    s4_codes = normalize_series(s4_df[S4_FIELDS["COMPANY_CODE"]])
    ecc_documents = ecc_df["Document Number"].apply(clean_string)
    s4_xref1 = s4_df[S4_FIELDS["REFERENCE_KEY_1"]].apply(clean_string)
    details = []

    for ecc_code in ecc_codes.replace("", pd.NA).dropna().unique():
        ecc_code = ecc_code.upper()
        s4_code = COMPANY_CODE_MAPPING.get(ecc_code)
        if not s4_code:
            details.append(make_detail(
                label=f"{ecc_code} (no mapping)", left_count=None, right_count=None,
                status="MAPPING_ERROR",
                message=f"No ECC -> S/4 company code mapping exists for {ecc_code}.",
                ecc_code=ecc_code, s4_code=None,
            ))
            continue

        s4_code = clean_string(s4_code).upper()
        ecc_mask = ecc_codes == ecc_code
        s4_mask = s4_codes == s4_code
        ecc_unique = int(ecc_documents[ecc_mask & (ecc_documents != "")].nunique())
        s4_unique = int(s4_xref1[s4_mask & (s4_xref1 != "")].nunique())

        details.append(make_detail(
            label=f"{ecc_code} -> {s4_code}", left_count=ecc_unique, right_count=s4_unique,
            message=("Unique document number count matches." if ecc_unique == s4_unique
                     else f"Unique document number count mismatch: ECC={ecc_unique}, S/4 XREF1={s4_unique}."),
            ecc_code=ecc_code, s4_code=s4_code,
            ecc_unique_document_count=ecc_unique, s4_unique_xref1_count=s4_unique,
        ))

    return make_check(
        check_name="Unique Document Number Count", details=details,
        passing_message="Unique ECC Document Number counts match unique S/4 XREF1 counts.",
        failing_message="Unique document number count validation failed.",
        company_code_wise=True,
    )


# ============================================================
# Check 6 - Document Cross-Reference
# ============================================================

def validate_document_cross_reference(ecc_df: pd.DataFrame, s4_df: pd.DataFrame) -> Dict[str, Any]:
    """Every ECC Document Number must appear as an S/4 XBLNR, and vice versa."""
    missing = first_missing_column(ecc_df, ["Document Number"])
    if missing:
        return missing_column_check("Document Cross-Reference", "ECC registry", missing)
    missing = first_missing_column(s4_df, [S4_FIELDS["DOCUMENT_NUMBER"]])
    if missing:
        return missing_column_check("Document Cross-Reference", "S/4 file", missing)

    ecc_docs = set(ecc_df["Document Number"].apply(clean_string)) - {""}
    s4_docs = set(s4_df[S4_FIELDS["DOCUMENT_NUMBER"]].apply(clean_string)) - {""}

    missing_in_s4 = sorted(ecc_docs - s4_docs)
    extra_in_s4 = sorted(s4_docs - ecc_docs)

    details = []
    for doc in missing_in_s4:
        details.append(make_detail(
            label=f"ECC {doc} not in S/4",
            left_count=1, right_count=0, status="FAIL",
            message=f"ECC document {doc} has no matching S/4 XBLNR.",
            direction="ECC_ONLY", document=doc,
        ))
    for doc in extra_in_s4:
        details.append(make_detail(
            label=f"S/4 {doc} not in ECC",
            left_count=0, right_count=1, status="FAIL",
            message=f"S/4 XBLNR {doc} has no matching ECC document.",
            direction="S4_ONLY", document=doc,
        ))

    if not details:
        details.append(make_detail(
            label=f"{len(ecc_docs)} documents matched",
            left_count=len(ecc_docs), right_count=len(s4_docs),
            status="PASS",
            message="Every ECC document appears as an S/4 XBLNR, and vice versa.",
        ))

    return make_check(
        check_name="Document Cross-Reference", details=details,
        passing_message="ECC and S/4 document sets match exactly.",
        failing_message="ECC and S/4 document sets differ.",
    )


# ============================================================
# Check 7 - Payment Terms Group Count
# ============================================================

def validate_payment_terms_group_counts(ecc_df: pd.DataFrame, s4_df: pd.DataFrame) -> Dict[str, Any]:
    """
    Compare payment-term counts grouped by the first letter of the mapped
    S/4 term. Groups: Net (N), Proxy (P), Discount (Z), E_payment_terms (E).
    """
    missing = first_missing_column(ecc_df, ["Terms of Payment"])
    if missing:
        return missing_column_check("Payment Terms Group Count", "ECC registry", missing)
    missing = first_missing_column(s4_df, [S4_FIELDS["PAYMENT_TERMS"]])
    if missing:
        return missing_column_check("Payment Terms Group Count", "S/4 file", missing)

    ecc_terms_raw = normalize_series(ecc_df["Terms of Payment"]).str.lstrip("0")
    s4_terms_raw = normalize_series(s4_df[S4_FIELDS["PAYMENT_TERMS"]])

    groups = {
        "N": {"label": "Net (N)", "ecc_count": 0, "s4_count": 0},
        "P": {"label": "Proxy (P)", "ecc_count": 0, "s4_count": 0},
        "Z": {"label": "Discount (Z)", "ecc_count": 0, "s4_count": 0},
        "E": {"label": "E_payment_terms (E)", "ecc_count": 0, "s4_count": 0},
    }

    for term in ecc_terms_raw:
        if term == "":
            continue
        mapped = _NORMALIZED_PAYMENT_MAPPING.get(term)
        if mapped is None:
            continue
        first_char = mapped[0] if mapped else ""
        if first_char in groups:
            groups[first_char]["ecc_count"] += 1

    NO_PAYMENT_TERMS_SENTINEL = "NO PAYMENT TERMS IN ECC"

    for term in s4_terms_raw:
        if term == "" or term == NO_PAYMENT_TERMS_SENTINEL:
            continue
        first_char = term[0] if term else ""
        if first_char in groups:
            groups[first_char]["s4_count"] += 1

    details = []
    failed_groups = []
    for prefix, data in groups.items():
        ecc_cnt = data["ecc_count"]
        s4_cnt = data["s4_count"]
        status = pass_fail(ecc_cnt == s4_cnt)
        if status == "FAIL":
            failed_groups.append(data["label"])
        details.append(make_detail(
            label=data["label"], left_count=ecc_cnt, right_count=s4_cnt, status=status,
            message=(f"{data['label']} counts match." if ecc_cnt == s4_cnt
                     else f"{data['label']} mismatch: ECC={ecc_cnt}, S/4={s4_cnt}."),
            prefix=prefix,
        ))

    total_ecc = sum(g["ecc_count"] for g in groups.values())
    total_s4 = sum(g["s4_count"] for g in groups.values())

    return make_check(
        check_name="Payment Terms Group Count",
        details=details,
        passing_message="All payment term groups (N, P, Z, E) have matching counts.",
        failing_message=("Payment term group validation failed. Mismatched groups: "
                         + ", ".join(failed_groups) + "."),
        ecc_total_count=total_ecc, s4_total_count=total_s4,
        difference=total_ecc - total_s4, groups=groups,
    )


# ============================================================
# Check 8 - Document Type
# ============================================================

def validate_document_type(ecc_df: pd.DataFrame, s4_df: pd.DataFrame) -> Dict[str, Any]:
    """Every S/4 row must carry BLART = 'UE'."""
    column = S4_FIELDS["DOCUMENT_TYPE"]
    missing = first_missing_column(s4_df, [column])
    if missing:
        return missing_column_check("Document Type", "S/4 file", missing)

    doc_types = normalize_series(s4_df[column])
    expected = mappings.get_s4_ar_document_type()

    total = len(doc_types)
    matching = int((doc_types == expected).sum())
    unexpected = total - matching

    details = [make_detail(
        label=f"BLART = {expected}",
        left_count=matching, right_count=total,
        status=pass_fail(unexpected == 0),
        message=(f"All {total} rows carry BLART={expected}."
                 if unexpected == 0
                 else f"{unexpected} of {total} rows carry an unexpected BLART."),
        expected=expected,
    )]

    return make_check(
        check_name="Document Type", details=details,
        passing_message=f"Every row carries BLART={expected}.",
        failing_message=f"One or more rows do not carry BLART={expected}.",
    )


# ============================================================
# Check 9 - Tax Code
# ============================================================

def validate_tax_code(ecc_df: pd.DataFrame, s4_df: pd.DataFrame) -> Dict[str, Any]:
    """Every S/4 row's MWSKZ must equal the tax code derived from its BUKRS."""
    missing = first_missing_column(s4_df, [S4_FIELDS["TAX_CODE"], S4_FIELDS["COMPANY_CODE"]])
    if missing:
        return missing_column_check("Tax Code", "S/4 file", missing)

    s4_codes = normalize_series(s4_df[S4_FIELDS["COMPANY_CODE"]])
    s4_tax = normalize_series(s4_df[S4_FIELDS["TAX_CODE"]])

    details = []
    for bukrs in s4_codes.replace("", pd.NA).dropna().unique():
        expected = get_tax_code(bukrs)
        mask = s4_codes == bukrs
        total = int(mask.sum())
        matching = int((s4_tax[mask] == expected).sum())
        details.append(make_detail(
            label=f"BUKRS={bukrs} | MWSKZ={expected}",
            left_count=matching, right_count=total,
            status=pass_fail(matching == total),
            message=(f"All {total} rows carry MWSKZ={expected}."
                     if matching == total
                     else f"{total - matching} of {total} rows do not carry MWSKZ={expected}."),
            s4_company_code=bukrs, expected_tax_code=expected,
        ))

    return make_check(
        check_name="Tax Code", details=details,
        passing_message="Every row's tax code matches the expected value for its company code.",
        failing_message="One or more rows carry the wrong tax code.",
    )


# ============================================================
# Check Registry
# ============================================================

CHECK_FUNCTIONS: List[Callable[[pd.DataFrame, pd.DataFrame], Dict[str, Any]]] = [
    validate_company_code_counts,
    validate_customer_distribution,
    validate_sign,
    validate_total_amount_reconciliation,
    validate_unique_document_number_counts,
    validate_document_cross_reference,
    validate_payment_terms_group_counts,
    validate_document_type,
    validate_tax_code,
]


# ============================================================
# Overall Validation
# ============================================================

def calculate_overall_status(checks: List[Dict[str, Any]]) -> str:
    return pass_fail(all(check.get("status") == "PASS" for check in checks))


def summarize_checks(checks: List[Dict[str, Any]]) -> Dict[str, int]:
    passed = sum(1 for check in checks if check["status"] == "PASS")
    return {
        "total_checks": len(checks),
        "passed": passed,
        "failed": len(checks) - passed,
    }


# ============================================================
# Main AR Validation Function
# ============================================================

def validate_ar_files(registry_file, filled_file) -> Dict[str, Any]:
    """
    Run every check in CHECK_FUNCTIONS against the ECC registry and the
    filled S/4 template. Returns a structured dict for FastAPI, the
    frontend, and the PDF report generator.
    """
    ecc_df = read_ecc_registry(registry_file)
    s4_df = read_s4_customer_open_items(filled_file)
    checks = [check_fn(ecc_df, s4_df) for check_fn in CHECK_FUNCTIONS]
    return {
        "process": "AR",
        "overall_status": calculate_overall_status(checks),
        "summary": summarize_checks(checks),
        "checks": checks,
    }