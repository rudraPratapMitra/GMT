# """Staging store for ECC_DATA / S4_DATA.

# Single-run mode: each folder holds at most one workbook. Writing clears
# the folder first so stale files cannot leak across runs.
# """
# import io
# import json
# from pathlib import Path

# import openpyxl

# from backend.config import ECC_DATA_DIR, S4_DATA_DIR
# from backend.services.ar_processor import clean_date

# # S/4 fields written as datetime.date objects by ar_processor.transform_record().
# # json.dumps(..., default=str) stringifies them to ISO ("2026-04-08") when the
# # sidecar is written; load_s4_rows() must reparse them back into date objects
# # or _SapJsonEncoder never converts them to OData's /Date(ms)/ format and SAP
# # rejects the push with CX_SY_CONVERSION_NO_DATE_TIME.
# S4_DATE_FIELDS = ("BLDAT", "ZFBDT")


# def _ensure(folder: Path) -> Path:
#     folder.mkdir(parents=True, exist_ok=True)
#     return folder


# def _clear(folder: Path) -> None:
#     _ensure(folder)
#     for entry in folder.iterdir():
#         if entry.is_file():
#             entry.unlink()


# def _latest(folder: Path) -> Path | None:
#     _ensure(folder)
#     files = [p for p in folder.iterdir()
#              if p.is_file() and p.suffix.lower() in (".xlsx", ".xls")]
#     if not files:
#         return None
#     files.sort(key=lambda p: p.stat().st_mtime, reverse=True)
#     return files[0]


# # --- ECC side --------------------------------------------------------------

# def save_ecc_workbook(buffer: io.BytesIO, filename: str) -> Path:
#     _clear(ECC_DATA_DIR)
#     path = _ensure(ECC_DATA_DIR) / filename
#     path.write_bytes(buffer.getvalue())
#     return path


# def latest_ecc_file() -> Path | None:
#     return _latest(ECC_DATA_DIR)


# def read_ecc_workbook(path: Path) -> list[dict]:
#     """Read the staged ECC Excel back into list-of-dicts for the processor."""
#     wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
#     ws = wb.active
#     rows = list(ws.iter_rows(values_only=True))
#     wb.close()
#     if not rows:
#         return []
#     headers = [str(h) if h is not None else "" for h in rows[0]]
#     return [dict(zip(headers, row)) for row in rows[1:]]


# # --- S/4 side --------------------------------------------------------------

# def save_s4_workbook(buffer: io.BytesIO, filename: str,
#                      s4_rows: list, warnings: list) -> Path:
#     """Write the LTMC workbook plus JSON sidecars holding the exact rows
#     the OData push will send. Sidecars avoid re-parsing the template layout
#     (row 5 = technical names, row 9 = headers, row 10 = data) at load time."""
#     _clear(S4_DATA_DIR)
#     folder = _ensure(S4_DATA_DIR)
#     path = folder / filename
#     path.write_bytes(buffer.getvalue())
#     (folder / "_s4_rows.json").write_text(json.dumps(s4_rows, default=str))
#     (folder / "_warnings.json").write_text(json.dumps(warnings, default=str))
#     return path


# def latest_s4_file() -> Path | None:
#     return _latest(S4_DATA_DIR)


# def load_s4_rows() -> list[dict]:
#     """Read the staged S/4 rows back from the JSON sidecar.

#     save_s4_workbook() writes dates as plain ISO strings (json.dumps has no
#     native date type), so BLDAT/ZFBDT are reparsed back into datetime.date
#     objects here -- otherwise push_to_s4()'s OData encoder never recognizes
#     them as dates and sends the raw string, which SAP rejects."""
#     sidecar = S4_DATA_DIR / "_s4_rows.json"
#     if not sidecar.exists():
#         return []

#     rows = json.loads(sidecar.read_text())
#     for row in rows:
#         for field in S4_DATE_FIELDS:
#             if row.get(field):
#                 row[field] = clean_date(row[field])
#     return rows


# def load_warnings() -> list[dict]:
#     sidecar = S4_DATA_DIR / "_warnings.json"
#     return json.loads(sidecar.read_text()) if sidecar.exists() else []


"""Staging store for ECC_DATA / S4_DATA.

Single-run mode: each folder holds at most one workbook. Writing clears
the folder first so stale files cannot leak across runs.
"""
import io
import json
from pathlib import Path

import openpyxl

from backend.config import ECC_DATA_DIR, S4_DATA_DIR
from backend.services.ar_processor import clean_date

# S/4 fields written as datetime.date objects by ar_processor.transform_record().
# json.dumps(..., default=str) stringifies them to ISO ("2026-04-08") when the
# sidecar is written; load_s4_rows() must reparse them back into date objects
# or _SapJsonEncoder never converts them to OData's /Date(ms)/ format and SAP
# rejects the push with CX_SY_CONVERSION_NO_DATE_TIME.
S4_DATE_FIELDS = ("BLDAT", "ZFBDT")


def _ensure(folder: Path) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def _clear(folder: Path) -> None:
    _ensure(folder)
    for entry in folder.iterdir():
        if entry.is_file():
            entry.unlink()


def _latest(folder: Path) -> Path | None:
    _ensure(folder)
    files = [p for p in folder.iterdir()
             if p.is_file() and p.suffix.lower() in (".xlsx", ".xls")]
    if not files:
        return None
    files.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    return files[0]


# --- ECC side --------------------------------------------------------------

def save_ecc_workbook(buffer: io.BytesIO, filename: str) -> Path:
    _clear(ECC_DATA_DIR)
    path = _ensure(ECC_DATA_DIR) / filename
    path.write_bytes(buffer.getvalue())
    return path


def latest_ecc_file() -> Path | None:
    return _latest(ECC_DATA_DIR)


def read_ecc_workbook(path: Path) -> list[dict]:
    """Read the staged ECC Excel back into list-of-dicts for the processor."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    headers = [str(h) if h is not None else "" for h in rows[0]]
    return [dict(zip(headers, row)) for row in rows[1:]]


def remove_ecc_rows(record_indices: list[int]) -> Path:
    """
    Deletes specific data rows from the currently staged ECC workbook, in
    place, and returns its path.

    `record_indices` are 1-based positions among the *data* rows, matching
    the numbering read_ecc_workbook() hands the processor (row 1 = the
    first data row, i.e. Excel row 2, since Excel row 1 is the header).
    This is exactly the "record" number ar_processor.py stamps on each
    warning/mismatch, so a caller can pass through the `record` values
    from a /process response's `currency_mismatches` unchanged.

    Edits the single staged file directly rather than going through
    save_ecc_workbook()'s clear-then-write, since there is already exactly
    one ECC file staged and its name/identity doesn't need to change.
    """
    path = latest_ecc_file()
    if path is None:
        raise FileNotFoundError("No ECC file staged.")
    if not record_indices:
        return path

    wb = openpyxl.load_workbook(path)
    ws = wb.active

    # Delete bottom-to-top so earlier deletions don't shift the row
    # numbers of rows still waiting to be deleted.
    excel_rows = sorted({idx + 1 for idx in record_indices}, reverse=True)
    for excel_row in excel_rows:
        ws.delete_rows(excel_row)

    wb.save(path)
    return path


# --- S/4 side --------------------------------------------------------------

def save_s4_workbook(buffer: io.BytesIO, filename: str,
                     s4_rows: list, warnings: list) -> Path:
    """Write the LTMC workbook plus JSON sidecars holding the exact rows
    the OData push will send. Sidecars avoid re-parsing the template layout
    (row 5 = technical names, row 9 = headers, row 10 = data) at load time."""
    _clear(S4_DATA_DIR)
    folder = _ensure(S4_DATA_DIR)
    path = folder / filename
    path.write_bytes(buffer.getvalue())
    (folder / "_s4_rows.json").write_text(json.dumps(s4_rows, default=str))
    (folder / "_warnings.json").write_text(json.dumps(warnings, default=str))
    return path


def latest_s4_file() -> Path | None:
    return _latest(S4_DATA_DIR)


def load_s4_rows() -> list[dict]:
    """Read the staged S/4 rows back from the JSON sidecar.

    save_s4_workbook() writes dates as plain ISO strings (json.dumps has no
    native date type), so BLDAT/ZFBDT are reparsed back into datetime.date
    objects here -- otherwise push_to_s4()'s OData encoder never recognizes
    them as dates and sends the raw string, which SAP rejects."""
    sidecar = S4_DATA_DIR / "_s4_rows.json"
    if not sidecar.exists():
        return []

    rows = json.loads(sidecar.read_text())
    for row in rows:
        for field in S4_DATE_FIELDS:
            if row.get(field):
                row[field] = clean_date(row[field])
    return rows


def load_warnings() -> list[dict]:
    sidecar = S4_DATA_DIR / "_warnings.json"
    return json.loads(sidecar.read_text()) if sidecar.exists() else []